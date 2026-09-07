"""Build a simple .docx from stored draft text for insertFileFromBase64."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Pt


def draft_to_docx_bytes(title: str, content: dict[str, Any] | None, fallback_text: str = "") -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if content and isinstance(content.get("content"), list):
        _walk(doc, content["content"], title_written=False)
    elif fallback_text:
        doc.add_heading(title, level=1)
        for para in fallback_text.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
    else:
        doc.add_heading(title or "Draft", level=1)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _node_text(node: dict[str, Any]) -> str:
    if isinstance(node.get("text"), str):
        return node["text"]
    children = node.get("content")
    if isinstance(children, list):
        return "".join(_node_text(c) for c in children if isinstance(c, dict))
    return ""


def _walk(doc: Document, blocks: list[Any], title_written: bool) -> None:
    for block in blocks:
        if not isinstance(block, dict):
            continue
        kind = block.get("type")
        text = _node_text(block).strip()
        if kind == "heading":
            level = 1
            attrs = block.get("attrs") or {}
            if isinstance(attrs.get("level"), int):
                level = max(1, min(3, attrs["level"]))
            if text:
                doc.add_heading(text, level=level)
        elif text:
            doc.add_paragraph(text)


AUTHOR = "HakiChain AI Contract Review"


def _revision_run(tag: str, text_tag: str, text: str, date: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    wrap = OxmlElement(tag)
    wrap.set(qn("w:author"), AUTHOR)
    wrap.set(qn("w:date"), date)
    wrap.set(qn("w:id"), "1")
    r = OxmlElement("w:r")
    t = OxmlElement(text_tag)
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    wrap.append(r)
    return wrap


def corrected_docx_bytes(
    document_text: str,
    accepted_redlines: list[dict[str, str]],
    tracked_changes: bool,
) -> bytes:
    """Apply accepted replacements. Skip redlines whose currentLanguage is missing."""
    from datetime import datetime, timezone

    from docx.oxml.ns import qn

    text = document_text
    applied: list[tuple[str, str]] = []
    for raw in accepted_redlines:
        current = raw.get("currentLanguage") or raw.get("current_language") or ""
        replacement = raw.get("replacementLanguage") or raw.get("replacement_language") or ""
        if not current or current not in text:
            continue
        text = text.replace(current, replacement, 1)
        applied.append((current, replacement))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    if tracked_changes and applied:
        remaining = document_text
        for current, replacement in applied:
            before, _, remaining = remaining.partition(current)
            if before:
                doc.add_paragraph(before)
            p = doc.add_paragraph()
            p._p.append(_revision_run("w:del", "w:delText", current, date))
            p._p.append(_revision_run("w:ins", "w:t", replacement, date))
        if remaining:
            doc.add_paragraph(remaining)
        settings = doc.settings.element
        track = settings.makeelement(qn("w:trackRevisions"), {})
        settings.append(track)
    else:
        for para in text.split("\n"):
            doc.add_paragraph(para)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
